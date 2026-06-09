#!/usr/bin/env python3
"""채용 면접 조편성 프로그램.

지원자 엑셀 파일을 읽어 지원부문별 4인 조를 만들고, 면접일자/시간을
랜덤 배정한 뒤 결과 엑셀 파일을 생성합니다.
"""

from __future__ import annotations

import argparse
import random
from collections import Counter
from dataclasses import dataclass
from datetime import date, time
from pathlib import Path
from typing import Any, Iterable

REQUIRED_COLUMNS = ["지원부문", "지원번호", "매칭직무", "출신학교", "성별"]
GROUP_SIZE = 4
DEFAULT_OUTPUT = "채용_면접_조편성_결과.xlsx"
DEFAULT_INTERVIEW_DATES = [date(2026, 6, 8), date(2026, 6, 9), date(2026, 6, 10)]
EXPECTED_GROUP_COUNTS = {
    "본사영업": 21,
    "본사지원": 11,
    "디지털금융": 12,
}
Row = dict[str, str]


@dataclass(frozen=True)
class Slot:
    interview_date: date
    interview_time: time

    @property
    def date_text(self) -> str:
        return self.interview_date.isoformat()

    @property
    def time_text(self) -> str:
        return self.interview_time.strftime("%H:%M")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="지원자 엑셀 파일을 기준으로 채용 면접 조편성 결과 엑셀을 생성합니다."
    )
    parser.add_argument("input", help="지원자 리스트 엑셀 파일 경로")
    parser.add_argument(
        "-o",
        "--output",
        default=DEFAULT_OUTPUT,
        help=f"결과 엑셀 파일 경로 (기본값: {DEFAULT_OUTPUT})",
    )
    parser.add_argument("--seed", type=int, default=None, help="재현 가능한 랜덤 배정을 위한 시드값")
    parser.add_argument(
        "--no-strict-counts",
        action="store_true",
        help="지원부문별 예상 인원/조 수 검증을 경고로만 처리합니다.",
    )
    return parser.parse_args()


def clean_cell(value: Any) -> str:
    if value is None:
        return "미기재"
    text = str(value).strip()
    return text if text else "미기재"


def read_applicants(input_path: Path) -> list[Row]:
    """엑셀을 읽고 필수 컬럼만 표준화해 반환합니다."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl이 필요합니다. `python -m pip install -r requirements.txt`를 먼저 실행하세요.") from exc

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        raise ValueError("입력 엑셀 파일이 비어 있습니다.") from exc

    headers = [str(value).strip() if value is not None else "" for value in header_row]
    header_indexes = {header: index for index, header in enumerate(headers)}
    missing_columns = [column for column in REQUIRED_COLUMNS if column not in header_indexes]
    if missing_columns:
        raise ValueError(
            "필수 컬럼이 없습니다: "
            + ", ".join(missing_columns)
            + "\n필수 컬럼: "
            + ", ".join(REQUIRED_COLUMNS)
        )

    applicants: list[Row] = []
    for row in rows:
        applicant = {column: clean_cell(row[header_indexes[column]]) for column in REQUIRED_COLUMNS}
        if applicant["지원번호"] != "미기재":
            applicants.append(applicant)
    return applicants


def calculate_group_count(applicant_count: int, division: str) -> int:
    if applicant_count % GROUP_SIZE != 0:
        raise ValueError(
            f"{division} 지원자 수가 {applicant_count}명입니다. "
            f"1개 조 {GROUP_SIZE}명 기준으로 나누어떨어져야 합니다."
        )
    return applicant_count // GROUP_SIZE


def validate_counts(applicants: list[Row], strict: bool) -> list[str]:
    """지원부문별 예상 인원과 조 수를 검증하고 경고 목록을 반환합니다."""
    warnings: list[str] = []
    actual_divisions = {row["지원부문"] for row in applicants}
    expected_divisions = set(EXPECTED_GROUP_COUNTS)
    unexpected = sorted(actual_divisions - expected_divisions)
    missing = sorted(expected_divisions - actual_divisions)

    if unexpected:
        warnings.append("예상하지 않은 지원부문: " + ", ".join(unexpected))
    if missing:
        warnings.append("데이터에 없는 예상 지원부문: " + ", ".join(missing))

    for division, expected_group_count in EXPECTED_GROUP_COUNTS.items():
        applicant_count = sum(1 for row in applicants if row["지원부문"] == division)
        if applicant_count == 0:
            continue
        actual_group_count = calculate_group_count(applicant_count, division)
        if actual_group_count != expected_group_count:
            warnings.append(
                f"{division}: 예상 {expected_group_count}개조/{expected_group_count * GROUP_SIZE}명, "
                f"실제 {actual_group_count}개조/{applicant_count}명"
            )

    if strict and warnings:
        raise ValueError("지원부문별 인원 검증 실패:\n- " + "\n- ".join(warnings))
    return warnings


def duplicate_count(values: Iterable[str]) -> int:
    counts = Counter(values)
    return sum(count - 1 for count in counts.values() if count > 1)


def balance_score(candidate: Row, group: list[Row], remaining: list[Row]) -> tuple[int, float]:
    """후보자를 현재 조에 넣을 때의 패널티 점수를 계산합니다.

    점수가 낮을수록 좋습니다. 매칭직무, 성별, 출신학교 순으로 큰 가중치를 둬서
    완전 균등이 어려울 때도 우선순위를 유지합니다.
    """
    values_after_add = {
        "매칭직무": [person["매칭직무"] for person in group] + [candidate["매칭직무"]],
        "성별": [person["성별"] for person in group] + [candidate["성별"]],
        "출신학교": [person["출신학교"] for person in group] + [candidate["출신학교"]],
    }
    duplicate_penalty = (
        duplicate_count(values_after_add["매칭직무"]) * 10_000
        + duplicate_count(values_after_add["성별"]) * 1_000
        + duplicate_count(values_after_add["출신학교"]) * 100
    )

    scarcity_bonus = 0.0
    for column, weight in (("매칭직무", 10), ("성별", 3), ("출신학교", 1)):
        value_count = sum(1 for row in remaining if row[column] == candidate[column])
        scarcity_bonus += weight / max(value_count, 1)

    completeness_penalty = (GROUP_SIZE - (len(group) + 1)) * 0.01
    return duplicate_penalty, completeness_penalty - scarcity_bonus


def create_balanced_groups(division_rows: list[Row], division: str, rng: random.Random) -> list[list[Row]]:
    group_count = calculate_group_count(len(division_rows), division)
    remaining = division_rows[:]
    rng.shuffle(remaining)
    groups: list[list[Row]] = [[] for _ in range(group_count)]

    while remaining:
        target_group = min(groups, key=lambda group: (len(group), rng.random()))
        candidate_indexes = list(range(len(remaining)))
        rng.shuffle(candidate_indexes)
        best_index = min(
            candidate_indexes,
            key=lambda index: balance_score(remaining[index], target_group, remaining),
        )
        target_group.append(remaining.pop(best_index))

    invalid_groups = [index + 1 for index, group in enumerate(groups) if len(group) != GROUP_SIZE]
    if invalid_groups:
        raise RuntimeError(f"{division}에서 4명이 아닌 조가 생성되었습니다: {invalid_groups}")
    return groups


def build_interview_slots(interview_dates: list[date]) -> list[Slot]:
    """09:00부터 50분 면접+10분 휴식 단위로 점심시간을 제외한 슬롯을 생성합니다."""
    daily_times = [
        time(9, 0),
        time(10, 0),
        time(11, 0),
        time(13, 0),
        time(14, 0),
        time(15, 0),
        time(16, 0),
        time(17, 0),
    ]
    return [Slot(interview_date, interview_time) for interview_date in interview_dates for interview_time in daily_times]


def assign_slots(group_keys_by_division: dict[str, list[str]], rng: random.Random) -> dict[str, Slot]:
    """지원부문별 독립 면접 트랙 안에서 겹치지 않게 시간 슬롯을 배정합니다."""
    base_slots = build_interview_slots(DEFAULT_INTERVIEW_DATES)
    slot_by_group: dict[str, Slot] = {}
    for division, group_keys in group_keys_by_division.items():
        if len(group_keys) > len(base_slots):
            raise ValueError(
                f"{division} 면접 슬롯이 부족합니다. 필요 {len(group_keys)}개, 가능 {len(base_slots)}개입니다. "
                "지원부문별 시간이 겹치지 않게 배정하려면 날짜 또는 하루 슬롯을 늘려야 합니다."
            )

        shuffled_slots = base_slots[:]
        rng.shuffle(shuffled_slots)
        shuffled_groups = group_keys[:]
        rng.shuffle(shuffled_groups)
        slot_by_group.update({group_key: slot for group_key, slot in zip(shuffled_groups, shuffled_slots)})
    return slot_by_group


def sorted_divisions(applicants: list[Row]) -> list[str]:
    divisions = {row["지원부문"] for row in applicants}
    return sorted(
        divisions,
        key=lambda value: (0, list(EXPECTED_GROUP_COUNTS).index(value)) if value in EXPECTED_GROUP_COUNTS else (1, value),
    )


def create_result(applicants: list[Row], seed: int | None = None) -> tuple[list[Row], list[Row], list[dict[str, object]]]:
    rng = random.Random(seed)
    summary_rows: list[Row] = []
    detail_rows: list[Row] = []
    validation_rows: list[dict[str, object]] = []
    group_keys_by_division: dict[str, list[str]] = {}

    for division in sorted_divisions(applicants):
        division_rows = [row for row in applicants if row["지원부문"] == division]
        groups = create_balanced_groups(division_rows, division, rng)

        for group_index, group in enumerate(groups, start=1):
            group_number = f"{division}-{group_index:02d}"
            group_keys_by_division.setdefault(division, []).append(group_number)
            validation_rows.append(
                {
                    "지원부문": division,
                    "조번호": group_number,
                    "인원": len(group),
                    "매칭직무_중복수": duplicate_count(row["매칭직무"] for row in group),
                    "성별_중복수": duplicate_count(row["성별"] for row in group),
                    "출신학교_중복수": duplicate_count(row["출신학교"] for row in group),
                    "매칭직무_구성": ", ".join(f"{key}:{value}" for key, value in Counter(row["매칭직무"] for row in group).items()),
                    "성별_구성": ", ".join(f"{key}:{value}" for key, value in Counter(row["성별"] for row in group).items()),
                    "출신학교_구성": ", ".join(f"{key}:{value}" for key, value in Counter(row["출신학교"] for row in group).items()),
                }
            )
            for row in group:
                summary_rows.append({"지원부문": division, "조번호": group_number, "지원번호": row["지원번호"]})
                detail_rows.append(
                    {
                        "지원부문": division,
                        "조번호": group_number,
                        "면접일자": "",
                        "면접시간": "",
                        "지원번호": row["지원번호"],
                        "매칭직무": row["매칭직무"],
                        "성별": row["성별"],
                        "출신학교": row["출신학교"],
                    }
                )

    slot_by_group = assign_slots(group_keys_by_division, rng)
    for row in detail_rows:
        slot = slot_by_group[row["조번호"]]
        row["면접일자"] = slot.date_text
        row["면접시간"] = slot.time_text

    return summary_rows, detail_rows, validation_rows


def write_sheet(worksheet: Any, headers: list[str], rows: list[dict[str, object]]) -> None:
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row.get(header, "") for header in headers])


def write_excel(output_path: Path, summary_rows: list[Row], detail_rows: list[Row], validation_rows: list[dict[str, object]]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise SystemExit("openpyxl이 필요합니다. `python -m pip install -r requirements.txt`를 먼저 실행하세요.") from exc

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "최종결과"
    write_sheet(summary_sheet, ["지원부문", "조번호", "지원번호"], summary_rows)

    detail_sheet = workbook.create_sheet("상세확인")
    write_sheet(detail_sheet, ["지원부문", "조번호", "면접일자", "면접시간", "지원번호", "매칭직무", "성별", "출신학교"], detail_rows)

    validation_sheet = workbook.create_sheet("중복검증")
    write_sheet(
        validation_sheet,
        ["지원부문", "조번호", "인원", "매칭직무_중복수", "성별_중복수", "출신학교_중복수", "매칭직무_구성", "성별_구성", "출신학교_구성"],
        validation_rows,
    )
    workbook.save(output_path)


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)

    applicants = read_applicants(input_path)
    warnings = validate_counts(applicants, strict=not args.no_strict_counts)
    summary_rows, detail_rows, validation_rows = create_result(applicants, seed=args.seed)
    write_excel(output_path, summary_rows, detail_rows, validation_rows)

    print(f"결과 파일 생성 완료: {output_path}")
    print(f"총 지원자 수: {len(summary_rows)}명")
    print(f"총 조 수: {len({row['조번호'] for row in summary_rows})}개")
    if warnings:
        print("검증 경고:")
        for warning in warnings:
            print(f"- {warning}")

    duplicate_totals = {
        "매칭직무_중복수": sum(int(row["매칭직무_중복수"]) for row in validation_rows),
        "성별_중복수": sum(int(row["성별_중복수"]) for row in validation_rows),
        "출신학교_중복수": sum(int(row["출신학교_중복수"]) for row in validation_rows),
    }
    print("중복 현황 합계:")
    for key, value in duplicate_totals.items():
        print(f"- {key}: {value}")


if __name__ == "__main__":
    main()
